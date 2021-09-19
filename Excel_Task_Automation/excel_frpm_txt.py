# Program that loading txt file to Excel sheet.

# Import libraries
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import *
from openpyxl.worksheet.table import Table,TableStyleInfo
#Open a txt file
text_file = open(r"C:\Udemy_Python_Masterclass\Excel_Task_Automation\employees.txt")
records = []

text_file.seek(0)
# Create list od lists rom txt file
for record in text_file.readlines():
    records.append(record.rstrip("\n").split(";"))

print (records)

workbook = Workbook()
# Save file xlsx
file_path = r"C:\Udemy_Python_Masterclass\Excel_Task_Automation\MyCompanyStaff.xlsx"
workbook.save(file_path)


print (workbook.sheetnames)
sheet = workbook['Sheet']
sheet.title = 'Employees'

for row in records:
    sheet.append(row)

table = Table(displayName = "Table", ref = "A1:G11")

print(openpyxl.worksheet.table.TABLESTYLES)

style = TableStyleInfo(name = "TableStyleMedium9", showRowStripes = True, showColumnStripes = True)

table.tableStyleInfo = style

sheet.add_table(table)

font = Font(color='FF000000', bold = True, italic = True)


for cell_no in range(2, 12):
	if int(sheet['G%s' % (cell_no)].value) > 55000:
		sheet['G%s' % (cell_no)].font = font

workbook.save(file_path)
text_file.close()
workbook.close()